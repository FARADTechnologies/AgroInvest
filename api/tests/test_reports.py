import io

from openpyxl import load_workbook


async def test_stats_for_staff(specialist):
    body = (await specialist.get("/api/v1/admin/stats")).json()
    assert body["applications_total"] >= 10
    assert body["competitions_open"] >= 3
    assert len(body["daily"]) == 30
    assert body["by_rayon"] and body["by_purpose"]


async def test_stats_hidden_from_entrepreneurs(entrepreneur):
    assert (await entrepreneur.get("/api/v1/admin/stats")).status_code == 403


async def test_public_stats_have_no_personal_data(client):
    response = await client.get("/api/v1/stats/public")
    body = response.json()
    assert body["applications_total"] > 0 and body["by_rayon"] and body["by_purpose"]
    assert "DRAFT" not in body["applications_by_status"]
    # Only published competitions; never drafts or ones waiting for approval.
    assert all(
        c["status"] in {"PUBLISHED", "CLOSED", "RESULTS_ANNOUNCED"} for c in body["competitions"]
    )
    # Aggregates only: no identifiers, e-mails, people or company names.
    text = response.text
    for private in ("1905832207", "5KQ7M2D", "agro-demo.test", "Nicat", "Yaşıl Vadi", "Leyla"):
        assert private not in text


async def test_excel_report_is_masked_and_readable(specialist):
    r = await specialist.get("/api/v1/admin/reports/applications.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Müraciətlər"]
    header = [c.value for c in ws[1]]
    assert header[0] == "Nömrə" and "FİN / VÖEN" in header
    ids = [row[5].value for row in ws.iter_rows(min_row=2)]
    assert ids and all("•" in v for v in ids)
    assert {"Xülasə", "Rayonlar üzrə"} <= set(wb.sheetnames)
